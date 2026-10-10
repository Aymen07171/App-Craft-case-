import os
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_6_designs_catalog():
    src_wb_path = os.path.abspath("CaseCraft_Master_30_Designs_Catalog.xlsx")
    if not os.path.isfile(src_wb_path):
        raise FileNotFoundError(f"Source catalog not found: {src_wb_path}")

    src_wb = openpyxl.load_workbook(src_wb_path)
    src_ws = src_wb["Master_Products"]

    # Target SKUs - exactly 1 per category
    target_skus = [
        "CASE-LOTUSS-001",
        "CASE-VOLCAN-001",
        "CASE-MOONBU-001",
        "CASE-PHOENI-001",
        "CASE-RONINS-001",
        "CASE-SEADRA-001"
    ]

    # Create new workbook
    new_wb = openpyxl.Workbook()
    new_ws = new_wb.active
    new_ws.title = "Master_Products"

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    ready_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    ready_font = Font(name="Calibri", size=11, bold=True, color="166534")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # Copy header
    headers = [cell.value for cell in src_ws[1]]
    new_ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = new_ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Map rows matching target SKUs
    target_row_map = {sku: None for sku in target_skus}
    for r in range(2, src_ws.max_row + 1):
        sku = src_ws.cell(row=r, column=1).value
        if sku in target_row_map and target_row_map[sku] is None:
            row_vals = [cell.value for cell in src_ws[r]]
            # Clean title to remove 'Vol. 01 (Celestial Sunburst)' if desired
            title_idx = 6  # 0-indexed column 7
            old_title = str(row_vals[title_idx])
            clean_title = old_title.replace(" Vol. 01 (Celestial Sunburst)", "").replace(" Vol. 01", "").strip()
            row_vals[title_idx] = clean_title
            target_row_map[sku] = row_vals

    for sku in target_skus:
        row_data = target_row_map[sku]
        if row_data:
            new_ws.append(row_data)

    # Style data rows
    for r_idx in range(2, new_ws.max_row + 1):
        for c_idx in range(1, len(headers) + 1):
            cell = new_ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)
            if c_idx in (1, 3, 5, 14, 15, 16, 17, 18):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Status cell green
            if c_idx == 18:
                cell.fill = ready_fill
                cell.font = ready_font

    # Auto-adjust column widths
    for col in new_ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        new_ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

    # Copy Sheet 2: Printify_34_Models
    if "Printify_34_Models" in src_wb.sheetnames:
        src_models_ws = src_wb["Printify_34_Models"]
        new_models_ws = new_wb.create_sheet(title="Printify_34_Models")
        for row in src_models_ws.iter_rows(values_only=True):
            new_models_ws.append(list(row))
        for col_idx in range(1, new_models_ws.max_column + 1):
            cell = new_models_ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in new_models_ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value or "")) for c in col)
            new_models_ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Copy Sheet 3: Variants_26_Manifest
    if "Variants_26_Manifest" in src_wb.sheetnames:
        src_var_ws = src_wb["Variants_26_Manifest"]
        new_var_ws = new_wb.create_sheet(title="Variants_26_Manifest")
        for row in src_var_ws.iter_rows(values_only=True):
            new_var_ws.append(list(row))
        for col_idx in range(1, new_var_ws.max_column + 1):
            cell = new_var_ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in new_var_ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(c.value or "")) for c in col)
            new_var_ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # Save to file targets
    targets = [
        "CaseCraft_Master_6_Designs_Catalog.xlsx",
        "public/CaseCraft_Master_6_Designs_Catalog.xlsx",
        "CaseCraft_Master_Catalog.xlsx",
        "public/CaseCraft_Master_Catalog.xlsx"
    ]
    for target in targets:
        full_p = os.path.abspath(target)
        new_wb.save(full_p)
        print(f"[OK] Saved 6-design catalog: {full_p}")

if __name__ == "__main__":
    generate_6_designs_catalog()
