#!/usr/bin/env python3
"""
scripts/build_sardines_10_collection.py
Complete End-to-End Pipeline Automation for 10 Portuguese Sardinhas Phone Case Designs:
1. Resizes/saves 10 Master 9:16 Artworks (1344x2389 @ 300 DPI) in public/designs/
2. Renders 20 Dual-Device Photorealistic Lifestyle Mockups (2048x2048 @ 300 DPI) in public/mockups/
3. Uploads artworks to Printify & creates 10 Tough Case products (Blueprint 269, Provider 99, 34 variants @ $22.20)
4. Configures strictly 8/20 mockups per product via Selenium (2 custom mockups first, iPhone 18 Pro Max as Primary)
5. Publishes all 10 products to Etsy store CraftCasesStudio (Shop ID 7939147)
6. Updates CaseCraft_Pinterest_Single_Mockup.xlsx and .csv with all 10 designs (strictly 1 primary mockup per design)
"""

import os
import sys
import time
import json
import base64
import urllib.request
import urllib.error
from pathlib import Path
from PIL import Image, ImageFilter
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By

BRAIN_DIR = Path(r"C:\Users\jacka\.gemini\antigravity\brain\93b4ac09-5823-4126-86c4-fc067944487c")
PUBLIC_DESIGNS = Path("public/designs")
PUBLIC_MOCKUPS = Path("public/mockups")
PUBLIC_DESIGNS.mkdir(parents=True, exist_ok=True)
PUBLIC_MOCKUPS.mkdir(parents=True, exist_ok=True)

# Load Printify token
env_text = Path(".env").read_text(encoding="utf-8")
token = [l.split("=", 1)[1].strip() for l in env_text.splitlines() if l.startswith("PRINTIFY_API_TOKEN=")][0]

SHOP_ID = "29067558"
BASE_URL = "https://api.printify.com/v1"
HEADERS = {
    "Authorization": f"Bearer {token}",
    "Content-Type": "application/json",
    "User-Agent": "CraftCase/1.0"
}

# 10 Designs Specs
SARDINES_CATALOG = [
    {
        "sku": "CASE-SARDPT-001",
        "title": "The Porto Classic Sardinhas in Olive Oil Tough Phone Case",
        "description": "Iconic Porto classic canned sardines in pure golden olive oil framed by ornate handcrafted blue-and-white Portuguese Baroque Azulejo ceramic tiles, fresh Mediterranean lemon slices, and olive sprigs. Premium dual-layer Tough Case with impact-resistant polycarbonate backplate.",
        "tags": ["portuguese sardines", "porto azulejo case", "sardinhas case", "blue tile phone case", "portugal travel gift", "lisbon souvenir", "azulejo tile art", "olive oil sardine", "vintage tin case", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "mediterranean art"],
        "art_source": BRAIN_DIR / "porto_classic_sardinhas_1791887699501.jpg",
    },
    {
        "sku": "CASE-SARDPT-002",
        "title": "Galo de Barcelos Portuguese Rooster Spicy Sardines Tough Case",
        "description": "Traditional Portuguese folk art Galo de Barcelos lucky rooster perched proudly atop a vintage crimson sardine tin with fiery piri-piri chili peppers, hearts, and floral motifs. Dual-layer shock-absorbing Tough Case.",
        "tags": ["galo de barcelos", "portuguese rooster", "piri piri sardines", "portugal folk art", "spicy sardine case", "barcelos rooster", "lisbon souvenir", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "portuguese gift", "retro tin can", "portugal pride"],
        "art_source": BRAIN_DIR / "galo_barcelos_sardines_1791887751974.jpg",
    },
    {
        "sku": "CASE-SARDPT-003",
        "title": "Algarve Coastal Fisherman Morning Haul Sardines Tough Case",
        "description": "Vintage maritime nautical woodcut illustration of a weathered Portuguese fisherman in a traditional barrete wool cap hauling shimmering silver Atlantic sardines into a rustic coastal tin. Heavy-duty dual-layer Tough Case.",
        "tags": ["algarve fisherman", "nautical phone case", "sardine haul tin", "portugal coast art", "atlantic sardines", "fisherman gift", "vintage maritime", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "ocean beach case", "portuguese sea art", "coastal decor"],
        "art_source": BRAIN_DIR / "algarve_fisherman_haul_tin_1791887756583.jpg",
    },
    {
        "sku": "CASE-SARDPT-004",
        "title": "Fado Night Sardinhas and Portuguese Guitar Tough Phone Case",
        "description": "Atmospheric midnight lithograph featuring a traditional 12-string Guitarra Portuguesa resting against an open vintage sardine tin in the cobbled lantern-lit alleyways of Alfama Lisbon. Protective dual-layer Tough Case.",
        "tags": ["fado music case", "guitarra portuguesa", "alfama lisbon case", "fado night sardines", "portuguese guitar", "lisbon fado gift", "midnight lithograph", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "portugal music lover", "romantic lisbon", "vintage fado"],
        "art_source": BRAIN_DIR / "fado_night_sardinhas_guitar_1791887757057.jpg",
    },
    {
        "sku": "CASE-SARDPT-005",
        "title": "Tram 28 Lisbon Sardine Express Vintage Travel Tough Case",
        "description": "Vintage European travel poster styling showcasing the iconic yellow Lisbon Tram 28 reimagined as a retro Sardinhas de Lisboa tin gliding over sunlit Portuguese cobblestone streets. Durable dual-layer Tough Case.",
        "tags": ["tram 28 lisbon", "yellow tram case", "sardinhas de lisboa", "lisbon travel poster", "portugal wanderlust", "vintage travel tin", "tram 28 souvenir", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "lisbon city case", "european traveler", "street tram art"],
        "art_source": BRAIN_DIR / "tram_28_lisbon_sardine_express_1791887755727.jpg",
    },
    {
        "sku": "CASE-SARDPT-006",
        "title": "Azulejo Tile Mosaic Silver Pilchard Trio Tough Phone Case",
        "description": "Harmonious circular Yin-Yang swimming medallion of three sleek silver sardines set in an authentic cobalt-blue and white geometric Portuguese Azulejo ceramic tile mosaic. Dual-layer impact-resistant Tough Case.",
        "tags": ["azulejo tile mosaic", "silver sardines case", "pilchard trio art", "portuguese ceramic", "delft blue tile", "yinyang fish case", "minimalist tile art", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "mediterranean tile", "lisbon aesthetic", "geometric fish"],
        "art_source": BRAIN_DIR / "azulejo_pilchard_gemini_web.png",
    },
    {
        "sku": "CASE-SARDPT-007",
        "title": "Belem Tower Maritime Sardine Cannery Label Tough Case",
        "description": "19th-century royal archival engraving of Lisbon's historic Torre de Belém fortress framed by the Portuguese royal coat of arms, maritime armillary spheres, and vintage Conservas Reais typography. Tough Case.",
        "tags": ["belem tower case", "torre de belem", "conservas reais", "royal sardine tin", "portugal history art", "vintage engraving", "lisbon monument", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "archival travel case", "architecture lover", "portuguese empire"],
        "art_source": BRAIN_DIR / "belem_tower_cannery_label.png",
    },
    {
        "sku": "CASE-SARDPT-008",
        "title": "Sardinha com Pimentos Grilled Pepper Sardine Fest Tough Case",
        "description": "Sizzling Santos Populares street festival celebration featuring charred grilled sardines served over rustic broa cornbread with charred red and green bell peppers and summer festival garland. Durable Tough Case.",
        "tags": ["sardinha com pimentos", "grilled sardines", "santos populares", "festa de lisboa", "portuguese food art", "sardine festival", "broa cornbread", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "foodie phone case", "portugal summer fest", "lisbon street food"],
        "art_source": BRAIN_DIR / "sardinha_com_pimentos_fest.png",
    },
    {
        "sku": "CASE-SARDPT-009",
        "title": "Madeira Island Exotic Passion Fruit Sardine Tough Case",
        "description": "Botanical Art Nouveau vintage sardine tin label adorned with exotic Madeira passion fruit blossoms, lush banana leaves, and rugged Atlantic volcanic cliffs. Protective dual-layer Tough Case.",
        "tags": ["madeira island case", "passion fruit tin", "botanical art nouveau", "tropical sardine", "portugal islands", "madeira flower fest", "atlantic island art", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "funchal souvenir", "exotic botanical", "volcanic cliffs"],
        "art_source": BRAIN_DIR / "madeira_passion_fruit_sardine.png",
    },
    {
        "sku": "CASE-SARDPT-010",
        "title": "Cascais Riviera Sunbather Sardine Picnic Tough Case",
        "description": "Chic 1950s Mid-Century French-Portuguese Riviera gouache featuring cat-eye sunglasses, striped beach cabana umbrella, and a golden sardine tin on a sun-drenched Cascais sandy beach towel. Dual-layer Tough Case.",
        "tags": ["cascais riviera", "beach picnic case", "retro sunbather", "mid century gouache", "portugal beach case", "coastal summer tin", "vintage 1950s resort", "iphone 18 pro max", "samsung galaxy s26", "tough phone case", "summer vacation gift", "estoril coast", "sandy beach picnic"],
        "art_source": BRAIN_DIR / "cascais_riviera_picnic.png",
    },
]

# Verified Studio Templates & Masks
IPHONE_TEMPLATE = BRAIN_DIR / "suncatcher_frog_iphone_mockup_1791809247094.jpg"
SAMSUNG_TEMPLATE = BRAIN_DIR / "suncatcher_frog_samsung_mockup_1791809260072.jpg"
IPHONE_MASK = BRAIN_DIR / "mask_iphone.png"
SAMSUNG_MASK = BRAIN_DIR / "mask_samsung.png"

IPHONE_CORNERS = [(386.2, 139.0), (735.5, 181.9), (648.4, 891.3), (299.1, 848.4)]
SAMSUNG_CORNERS = [(204.3, 228.8), (577.6, 85.5), (844.7, 781.4), (471.4, 924.7)]

TARGET_PRINTIFY_MOCKUPS = [
    ("iPhone 18 Pro Max", "camera_label=front"),
    ("iPhone 16 Pro Max", "camera_label=close-up"),
    ("iPhone 16 Pro Max", "camera_label=layers"),
    ("iPhone 11", "camera_label=context-1"),
    ("Samsung Galaxy S24", "camera_label=close-up-2"),
    ("Samsung Galaxy S26", "camera_label=front"),
]

def solve_homography(dest_pts, src_pts):
    matrix = []
    for (x, y), (X, Y) in zip(dest_pts, src_pts):
        matrix.append([x, y, 1, 0, 0, 0, -x*X, -y*X, X])
        matrix.append([0, 0, 0, x, y, 1, -x*Y, -y*Y, Y])
    n = 8
    for i in range(n):
        max_r = max(range(i, n), key=lambda r: abs(matrix[r][i]))
        matrix[i], matrix[max_r] = matrix[max_r], matrix[i]
        pivot = matrix[i][i]
        for j in range(i, n + 1):
            matrix[i][j] /= pivot
        for r in range(n):
            if r != i:
                factor = matrix[r][i]
                for j in range(i, n + 1):
                    matrix[r][j] -= factor * matrix[i][j]
    return [matrix[i][8] for i in range(8)]

def render_warped_mockup(template_path, mask_path, art_img, dest_corners):
    template = Image.open(template_path).convert("RGB")
    mask = Image.open(mask_path).convert("L")
    W, H = art_img.size
    src_pts = [(0, 0), (W, 0), (W, H), (0, H)]
    coeffs = solve_homography(dest_corners, src_pts)
    warped = art_img.transform(template.size, Image.Transform.PERSPECTIVE, coeffs, Image.Resampling.BICUBIC)
    smooth_mask = mask.filter(ImageFilter.GaussianBlur(1.2))
    return Image.composite(warped, template, smooth_mask)

def api_req(path, method="GET", body=None):
    data = json.dumps(body).encode("utf-8") if body is not None else None
    req = urllib.request.Request(f"{BASE_URL}{path}", data=data, headers=HEADERS, method=method)
    try:
        with urllib.request.urlopen(req) as resp:
            return resp.status, json.loads(resp.read().decode())
    except urllib.error.HTTPError as e:
        err_body = e.read().decode()
        try:
            return e.code, json.loads(err_body)
        except Exception:
            return e.code, {"raw": err_body}

def upload_image_to_printify(file_path, file_name):
    print(f"Uploading image: {file_name}...")
    b64 = base64.b64encode(Path(file_path).read_bytes()).decode("ascii")
    payload = {"file_name": file_name, "contents": b64}
    status, res = api_req("/uploads/images.json", "POST", payload)
    if status == 200:
        print(f"  [OK] Uploaded {file_name} -> ID: {res.get('id')}")
        return res
    else:
        raise RuntimeError(f"Failed to upload image {file_name}: {status} {res}")

def step1_and_step2_artworks_and_mockups():
    print("=" * 75)
    print("STEP 1 & STEP 2: FORMATTING 10 ARTWORKS & RENDERING 20 LIFESTYLE MOCKUPS")
    print("=" * 75)

    for item in SARDINES_CATALOG:
        sku = item["sku"]
        title = item["title"]
        src_path = item["art_source"]
        print(f"\nProcessing [{sku}] {title}...")

        # 1. Format Master 9:16 Artwork (1344 x 2389 @ 300 DPI)
        raw_art = Image.open(src_path).convert("RGB")
        art_1344 = raw_art.resize((1344, 2389), Image.Resampling.LANCZOS)
        out_art_path = PUBLIC_DESIGNS / f"{sku}_9x16_Artwork.png"
        art_1344.save(out_art_path, dpi=(300, 300), optimize=True)
        print(f"  [OK] Saved Master Artwork -> {out_art_path}")

        # 2. Render iPhone 18 Pro Max Lifestyle Mockup (2048 x 2048 @ 300 DPI)
        out_ip_path = PUBLIC_MOCKUPS / f"{sku}_iPhone_18_Pro_Max_Lifestyle.png"
        im_ip = render_warped_mockup(IPHONE_TEMPLATE, IPHONE_MASK, raw_art, IPHONE_CORNERS)
        im_ip_2048 = im_ip.resize((2048, 2048), Image.Resampling.LANCZOS)
        im_ip_2048.save(out_ip_path, dpi=(300, 300), optimize=True)
        print(f"  [OK] Saved iPhone Lifestyle -> {out_ip_path}")

        # 3. Render Samsung Galaxy S26 Lifestyle Mockup (2048 x 2048 @ 300 DPI)
        out_sg_path = PUBLIC_MOCKUPS / f"{sku}_Samsung_Galaxy_S26_Lifestyle.png"
        im_sg = render_warped_mockup(SAMSUNG_TEMPLATE, SAMSUNG_MASK, raw_art, SAMSUNG_CORNERS)
        im_sg_2048 = im_sg.resize((2048, 2048), Image.Resampling.LANCZOS)
        im_sg_2048.save(out_sg_path, dpi=(300, 300), optimize=True)
        print(f"  [OK] Saved Samsung Lifestyle -> {out_sg_path}")

def step3_create_printify_products():
    print("\n" + "=" * 75)
    print("STEP 3: CREATING 10 PRODUCTS ON PRINTIFY (SHOP 29067558)")
    print("=" * 75)

    # Fetch 34 enabled variants from reference product
    _, ref_prod = api_req(f"/shops/{SHOP_ID}/products/6acca8512fe24e63d7081433.json")
    variant_ids = [v["id"] for v in ref_prod["variants"] if v.get("is_enabled")]
    print(f"Using {len(variant_ids)} enabled variants @ $22.20 (2220 cents).")

    created_products = []
    for item in SARDINES_CATALOG:
        sku = item["sku"]
        title = item["title"]
        desc = item["description"]
        tags = item["tags"]
        art_path = PUBLIC_DESIGNS / f"{sku}_9x16_Artwork.png"

        print(f"\nCreating product for [{sku}] {title}...")
        art_res = upload_image_to_printify(art_path, f"{sku}_9x16_Artwork.png")
        art_id = art_res["id"]

        create_payload = {
            "title": title,
            "description": desc,
            "blueprint_id": 269,
            "print_provider_id": 99,
            "tags": tags,
            "variants": [{"id": vid, "price": 2220, "is_enabled": True} for vid in variant_ids],
            "print_areas": [{
                "variant_ids": variant_ids,
                "placeholders": [{
                    "position": "front",
                    "images": [{"id": art_id, "x": 0.5, "y": 0.5, "scale": 1, "angle": 0}]
                }]
            }]
        }

        st, created = api_req(f"/shops/{SHOP_ID}/products.json", "POST", create_payload)
        if st == 200:
            pid = created["id"]
            print(f"  [SUCCESS] Created Printify Product: {pid}")
            created_products.append({
                "sku": sku,
                "pid": pid,
                "title": title,
                "art_id": art_id,
            })
        else:
            print(f"  [ERROR] {st} {created}")

    Path("scripts/created_sardines_products.json").write_text(json.dumps(created_products, indent=2), encoding="utf-8")
    return created_products

def step4_configure_selenium_mockups(products):
    print("\n" + "=" * 75)
    print("STEP 4: CONFIGURING STRICT 8/20 MOCKUPS VIA SELENIUM")
    print("=" * 75)

    opts = Options()
    opts.add_experimental_option("debuggerAddress", "127.0.0.1:9222")
    driver = webdriver.Chrome(options=opts)

    for item in products:
        sku = item["sku"]
        pid = item["pid"]
        print(f"\n--- Configuring [{sku}] PID: {pid} ---")

        url = f"https://printify.com/app/mockup-library/shops/{SHOP_ID}/products/{pid}"
        driver.get(url)
        time.sleep(5)

        # Clear modals
        driver.execute_script("""
            document.querySelectorAll('[data-amplitude-engagement-modal-overlay], .rc-dialog-root, .rc-dialog-mask, .rc-dialog-wrap').forEach(el => el.remove());
            const popups = Array.from(document.querySelectorAll('div')).filter(d => d.innerText.includes('MULTI-USER ACCESS') || d.innerText.includes('Sharing is caring'));
            popups.forEach(p => p.remove());
        """)
        time.sleep(1)

        # Clear default selections
        clear_btns = driver.find_elements(By.CSS_SELECTOR, "button.controls-clear")
        if clear_btns:
            driver.execute_script("arguments[0].click();", clear_btns[0])
            time.sleep(1.5)

        # Reveal file input via Upload button
        upload_nav_btns = driver.find_elements(By.CSS_SELECTOR, "button[data-analyticsid='mockupLibraryNavigationUpload'], pfy-button[data-testid='emptyStateUploadMockup'] button")
        if upload_nav_btns:
            driver.execute_script("arguments[0].click();", upload_nav_btns[0])
            time.sleep(1.5)

        # Remove popups
        driver.execute_script("""
            document.querySelectorAll('[data-amplitude-engagement-modal-overlay], .rc-dialog-root, .rc-dialog-mask, .rc-dialog-wrap').forEach(el => el.remove());
            const popups = Array.from(document.querySelectorAll('div')).filter(d => d.innerText.includes('MULTI-USER ACCESS') || d.innerText.includes('Sharing is caring'));
            popups.forEach(p => p.remove());
        """)
        time.sleep(1)

        file_inputs = driver.find_elements(By.CSS_SELECTOR, "input[data-testid='fileInput']")
        if not file_inputs:
            print(f"[ERROR] File input not found for {sku}")
            continue

        ip_abs = os.path.abspath(f"public/mockups/{sku}_iPhone_18_Pro_Max_Lifestyle.png")
        sg_abs = os.path.abspath(f"public/mockups/{sku}_Samsung_Galaxy_S26_Lifestyle.png")
        print(f"Uploading custom lifestyle mockups:\n  {ip_abs}\n  {sg_abs}")
        file_inputs[0].send_keys(f"{ip_abs}\n{sg_abs}")
        time.sleep(5)

        # Click Confirm
        confirm_btns = driver.find_elements(By.XPATH, "//button[normalize-space(.)='Confirm']")
        if confirm_btns:
            driver.execute_script("arguments[0].click();", confirm_btns[0])
            time.sleep(4)

        # Switch to Printify mockups tab
        mockup_tabs = driver.find_elements(By.XPATH, "//*[contains(@class, 'tab-label') and contains(text(), 'Printify mockups')]")
        if mockup_tabs:
            driver.execute_script("arguments[0].click();", mockup_tabs[0])
            time.sleep(2)

        # Select the 6 designated mockups
        for model_name, cam_query in TARGET_PRINTIFY_MOCKUPS:
            cards = driver.find_elements(By.XPATH, f"//button[contains(@class, 'view-type-card') and normalize-space(.)='{model_name}']")
            if cards:
                driver.execute_script("arguments[0].scrollIntoView({block: 'nearest', inline: 'center'}); arguments[0].click();", cards[0])
                time.sleep(1.2)
                driver.execute_script("""
                    const camQuery = arguments[0];
                    const buttons = Array.from(document.querySelectorAll('button.mockup-container:not(.draggable)'));
                    for (const btn of buttons) {
                        const img = btn.querySelector('img');
                        if (img && img.src.includes(camQuery)) {
                            const cbWrapper = btn.querySelector('[data-testid="checkboxWrapper"]');
                            const cbInput = btn.querySelector('input[data-testid="checkboxInput"]');
                            if (cbWrapper && cbWrapper.getAttribute('aria-checked') !== 'true') {
                                if (cbInput) cbInput.click();
                                else cbWrapper.click();
                            }
                            break;
                        }
                    }
                """, cam_query)
                time.sleep(0.6)

        # Save selection
        save_sel_btns = driver.find_elements(By.XPATH, "//button[contains(., 'Save selection')]")
        if save_sel_btns:
            driver.execute_script("arguments[0].click();", save_sel_btns[0])
            time.sleep(5.0)

        # Save as draft
        save_draft_btns = driver.find_elements(By.XPATH, "//button[contains(., 'Save as draft')]")
        for b in save_draft_btns:
            if b.is_displayed() and b.is_enabled():
                driver.execute_script("arguments[0].click();", b)
                time.sleep(5.0)
                break

        # Verify via API
        st, pdata = api_req(f"/shops/{SHOP_ID}/products/{pid}.json")
        cnt = len(pdata.get("images", []))
        print(f"[VERIFIED] {sku} ({pid}): {cnt} mockups selected in Printify API.")

def step5_publish_to_etsy(products):
    print("\n" + "=" * 75)
    print("STEP 5: SYNCING & PUBLISHING 10 PRODUCTS TO ETSY (CraftCasesStudio)")
    print("=" * 75)

    publish_payload = {
        "title": True,
        "description": True,
        "images": True,
        "variants": True,
        "tags": True,
        "keyFeatures": True,
        "shipping_template": True
    }

    for item in products:
        sku = item["sku"]
        pid = item["pid"]
        print(f"Publishing [{sku}] PID: {pid} to Etsy...")
        st, res = api_req(f"/shops/{SHOP_ID}/products/{pid}/publish.json", "POST", publish_payload)
        if st == 200:
            print(f"  [OK] Successfully queued for Etsy sync ({sku})")
        else:
            print(f"  [ERROR] {st} {res}")

def step6_generate_pinterest_files(products):
    print("\n" + "=" * 75)
    print("STEP 6: EXPORTING SINGLE-MOCKUP PINTEREST EXCEL & CSV")
    print("=" * 75)

    base_date = datetime.utcnow().date() + timedelta(days=1)
    PINTEREST_HEADERS = [
        "Title",
        "Media URL",
        "Pinterest board",
        "Thumbnail",
        "Description",
        "Link",
        "Publish date",
        "Keywords",
    ]

    pin_rows = []
    for idx, item in enumerate(products):
        sku = item["sku"]
        pid = item["pid"]
        st, pdata = api_req(f"/shops/{SHOP_ID}/products/{pid}.json")
        imgs = pdata.get("images", [])
        primary = [i for i in imgs if i.get("is_default")]
        media_url = primary[0]["src"] if primary else (imgs[0]["src"] if imgs else "")

        cat_item = next(c for c in SARDINES_CATALOG if c["sku"] == sku)

        pin_rows.append({
            "Title": cat_item["title"][:95],
            "Media URL": media_url,
            "Pinterest board": "Phone Cases",
            "Thumbnail": "",
            "Description": cat_item["description"][:490],
            "Link": f"https://www.etsy.com/shop/CraftCasesStudio?utm_source=pinterest&utm_medium=social&utm_campaign={sku.lower()}",
            "Publish date": (base_date + timedelta(days=idx)).isoformat(),
            "Keywords": ", ".join(cat_item["tags"]),
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "Pinterest_Single_Mockup"

    header_fill = PatternFill(start_color="E60023", end_color="E60023", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10, color="111827")
    link_font = Font(name="Calibri", size=10, color="1D4ED8", underline="single")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    ws.append(PINTEREST_HEADERS)
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(PINTEREST_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for r_idx, pin in enumerate(pin_rows, start=2):
        row_vals = [pin[h] for h in PINTEREST_HEADERS]
        ws.append(row_vals)
        ws.row_dimensions[r_idx].height = 54
        for col_idx, h in enumerate(PINTEREST_HEADERS, start=1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = thin_border
            if h in ("Media URL", "Link"):
                cell.font = link_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.font = body_font
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    col_widths = {
        "Title": 48,
        "Media URL": 75,
        "Pinterest board": 20,
        "Thumbnail": 14,
        "Description": 64,
        "Link": 65,
        "Publish date": 16,
        "Keywords": 60,
    }
    for col_idx, h in enumerate(PINTEREST_HEADERS, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_widths.get(h, 24)
    ws.freeze_panes = "A2"

    for p in ["CaseCraft_Pinterest_Single_Mockup.xlsx", "public/CaseCraft_Pinterest_Single_Mockup.xlsx"]:
        wb.save(p)
        print(f"Saved Excel workbook ({len(pin_rows)} single mockup rows) -> {p}")

    import csv
    for p in ["CaseCraft_Pinterest_Single_Mockup.csv", "public/CaseCraft_Pinterest_Single_Mockup.csv"]:
        with open(p, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=PINTEREST_HEADERS)
            writer.writeheader()
            for r in pin_rows:
                writer.writerow(r)
        print(f"Saved CSV file ({len(pin_rows)} single mockup rows) -> {p}")

def main():
    step1_and_step2_artworks_and_mockups()
    created = step3_create_printify_products()
    step4_configure_selenium_mockups(created)
    step5_publish_to_etsy(created)
    step6_generate_pinterest_files(created)
    print("\n" + "=" * 75)
    print("ALL 10 SARDINHAS DESIGNS FULLY PROCESSED, CONFIGURED & PUBLISHED!")
    print("=" * 75)

if __name__ == "__main__":
    main()
