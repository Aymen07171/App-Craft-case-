import csv
import os
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Official Pinterest Bulk Upload Columns
PINTEREST_OFFICIAL_HEADERS = [
    "Title",
    "Media URL",
    "Pinterest board",
    "Thumbnail",
    "Description",
    "Link",
    "Publish date",
    "Keywords",
]

# Single Mockup Row (Stained Glass Vitrail 70s Frog & Lava Lamp — Primary Custom Lifestyle Mockup iPhone 18 Pro Max)
publish_date_str = (datetime.utcnow().date() + timedelta(days=1)).isoformat()

SINGLE_MOCKUP_PIN = {
    "Title": "Stained Glass Vitrail 70s Frog & Lava Lamp Tough Phone Case | Groovy Retro Art",
    "Media URL": "https://d123s6f1z9g2wk.cloudfront.net/files/2026/10/20261012100924-1f1c6250-00e7-6b38-b2b4-5e9ef2da5bd5.png",
    "Pinterest board": "Phone Cases",
    "Thumbnail": "",
    "Description": (
        "Groovy 1970s stained-glass vitrail tree frog wearing retro amber sunglasses beside a glowing molten lava lamp, "
        "surrounded by psychedelic mushrooms and tropical monstera leaves. Dual-layer impact-resistant polycarbonate outer shell "
        "with shock-absorbing TPU inner liner. Fits 34 iPhone & Samsung Galaxy models."
    ),
    "Link": "https://www.etsy.com/shop/CraftCasesStudio?utm_source=pinterest&utm_medium=social&utm_campaign=stained_glass_frog_001",
    "Publish date": publish_date_str,
    "Keywords": (
        "stained glass case, 70s frog phone case, lava lamp case, groovy retro frog, vitrail phone case, "
        "mushroom frog case, iphone 18 pro max, samsung galaxy s26, tough phone case, boho hippie case, "
        "psychedelic frog, cottagecore frog, gift for frog lover"
    ),
}


def build_single_mockup_excel(output_paths: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pinterest_Single_Mockup"

    header_fill = PatternFill(start_color="E60023", end_color="E60023", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=10, color="111827")
    link_font = Font(name="Calibri", size=10, color="1D4ED8", underline="single")
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    # Row 1: Official Pinterest Bulk Upload Headers
    ws.append(PINTEREST_OFFICIAL_HEADERS)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(PINTEREST_OFFICIAL_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Row 2: Strictly 1 Mockup Row
    row_values = [SINGLE_MOCKUP_PIN[h] for h in PINTEREST_OFFICIAL_HEADERS]
    ws.append(row_values)
    ws.row_dimensions[2].height = 54

    for col_idx, header_name in enumerate(PINTEREST_OFFICIAL_HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx)
        cell.border = thin_border
        if header_name in ("Media URL", "Link"):
            cell.font = link_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
        else:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    col_widths = {
        "Title": 48,
        "Media URL": 68,
        "Pinterest board": 20,
        "Thumbnail": 14,
        "Description": 64,
        "Link": 60,
        "Publish date": 16,
        "Keywords": 58,
    }
    for col_idx, header_name in enumerate(PINTEREST_OFFICIAL_HEADERS, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_widths.get(header_name, 24)

    ws.freeze_panes = "A2"

    for path in output_paths:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        wb.save(path)
        print(f"Saved Excel workbook (1 mockup row) -> {path}")


def build_single_mockup_csv(output_paths: list[str]) -> None:
    for path in output_paths:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        with open(path, "w", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=PINTEREST_OFFICIAL_HEADERS, quoting=csv.QUOTE_MINIMAL)
            writer.writeheader()
            writer.writerow(SINGLE_MOCKUP_PIN)
        print(f"Saved CSV file (1 mockup row) -> {path}")


if __name__ == "__main__":
    build_single_mockup_excel([
        "CaseCraft_Pinterest_Single_Mockup.xlsx",
        "public/CaseCraft_Pinterest_Single_Mockup.xlsx",
    ])
    build_single_mockup_csv([
        "CaseCraft_Pinterest_Single_Mockup.csv",
        "public/CaseCraft_Pinterest_Single_Mockup.csv",
    ])
